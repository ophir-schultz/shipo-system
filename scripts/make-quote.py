#!/usr/bin/env python3
"""Produce a finished, sendable price quote for one customer.

    python3 scripts/make-quote.py "Acme Supplements" \\
        "19102,2nd,3" "30303,2nd,12,14x12x10" "90001,3day,25" "60601,ground,8"

Each box is  ZIP,service,weight[,LxWxH]  where service is matched loosely:
2nd / 3day / ground / nda ... whatever ShipStation returned for the account.

Writes  out/Shipo-Quote-<customer>.xlsx  and prints the same quote to the screen.

Every number traces to a live ShipStation rate in out/ups-rates-raw.json. The
zone comes from the official UPS chart for origin 198. Nothing is estimated: a
ZIP the chart does not list, or a weight/zone ShipStation did not return, is
refused outright rather than filled in with a plausible-looking number.
"""
import importlib.util
import math
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("bpl", os.path.join(HERE, "build-price-list.py"))
bpl = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bpl)

DIVISOR = 166          # CONFIRMED by Ophir 2026-09-03
OUT = os.path.join(HERE, "out")

ARIAL = "Arial"
BOLD = Font(name=ARIAL, size=10, bold=True)
BODY = Font(name=ARIAL, size=10)
BLUE = Font(name=ARIAL, size=10, color="0000FF")
HDRFILL = PatternFill("solid", fgColor="DDDDDD")
MONEY = '$#,##0.00;($#,##0.00);-'
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def markup_for(lb):
    for lo, hi, mk in bpl.LADDER:
        if lo <= lb <= hi:
            return mk
    return None


def main():
    if len(sys.argv) < 3:
        print(__doc__.strip())
        sys.exit(1)
    customer = sys.argv[1]

    data = bpl.load()
    rows, _ = bpl.apply_zone_chart(data["rows"])
    carrier = bpl.pick_baseline(rows)
    services = bpl.discover_services(rows, carrier)
    chart = bpl.load_zone_chart()

    lines = []
    for spec_str in sys.argv[2:]:
        parts = [p.strip() for p in spec_str.split(",")]
        if len(parts) < 3:
            print(f"BAD LINE '{spec_str}' — need ZIP,service,weight[,LxWxH]")
            sys.exit(2)
        zip_code, svc_in, wt = parts[0], parts[1], float(parts[2])
        dims = None
        if len(parts) > 3 and parts[3]:
            d = parts[3].lower().split("x")
            if len(d) != 3:
                print(f"BAD DIMS '{parts[3]}' — use LxWxH, e.g. 14x12x10")
                sys.exit(2)
            dims = [float(x) for x in d]

        prefix = "".join(ch for ch in zip_code if ch.isdigit())[:3].zfill(3)
        zone = chart.get(prefix)
        if zone is None:
            print(f"REFUSED: ZIP {zip_code} is not on the UPS origin-198 chart "
                  f"(APO/FPO, AK, HI). Quote it in ShipStation directly.")
            sys.exit(2)

        # Shorthand you can actually type at 6pm without checking the ® spelling.
        alias = {"2nd": "2nd day air", "2day": "2nd day air", "2": "2nd day air",
                 "3day": "3 day select", "3": "3 day select", "3 day": "3 day select",
                 "ground": "ground", "gs": "ground saver",
                 "nda": "next day air", "overnight": "next day air"}
        svc = bpl.resolve_service(services, alias.get(svc_in.lower(), svc_in))
        if not svc:
            print(f"REFUSED: no service matching '{svc_in}'. Available:")
            for s in services:
                print("   ", s)
            sys.exit(2)

        dim_lb = math.ceil(dims[0] * dims[1] * dims[2] / DIVISOR) if dims else 0
        billable = max(math.ceil(wt), dim_lb)
        mk = markup_for(billable)
        cost = bpl.service_costs(rows, svc, carrier).get((zone, billable))
        if cost is None or mk is None:
            print(f"REFUSED: no live rate for zone {zone} at {billable} lb on {svc}. "
                  f"Nothing was invented to fill the hole.")
            sys.exit(2)
        lines.append({
            "zip": zip_code, "svc": svc, "zone": zone, "actual": wt,
            "dims": dims, "dim_lb": dim_lb, "billable": billable,
            "cost": cost, "markup": mk, "price": round(cost + mk, 2),
        })

    # ---- screen version ----
    w = 96
    print("=" * w)
    print("SHIPO LLC — SHIPPING QUOTE")
    print("310 Cornell Dr, Wilmington, DE 19801  ·  302-442-2343  ·  shipousa.com")
    print("=" * w)
    print(f"Customer: {customer}")
    print(f"Quote date: {str(data.get('pulledAt',''))[:10]}   ·   valid 14 days")
    print(f"Rates: live UPS via ShipStation, account '{carrier}', residential, from ZIP "
          f"{data.get('fromZip','')}")
    print("-" * w)
    print(f"{'#':<3}{'To ZIP':<9}{'Service':<26}{'Zone':>5}{'Actual':>8}{'Billable':>10}{'PRICE':>12}")
    print("-" * w)
    for i, l in enumerate(lines, 1):
        print(f"{i:<3}{l['zip']:<9}{l['svc'][:25]:<26}{l['zone']:>5}"
              f"{l['actual']:>7g}{'lb':<1}{l['billable']:>9}{'':<1}"
              f"{'$%.2f' % l['price']:>12}")
    total = sum(l["price"] for l in lines)
    print("-" * w)
    print(f"{'TOTAL':<51}{len(lines)} shipment(s){'$%.2f' % total:>18}")
    print("=" * w)
    for l in lines:
        if l["dims"] and l["dim_lb"] > math.ceil(l["actual"]):
            print(f"  note: {l['zip']} bills at {l['billable']} lb "
                  f"(dimensional {l['dims'][0]:g}x{l['dims'][1]:g}x{l['dims'][2]:g}/{DIVISOR}), "
                  f"not its {l['actual']:g} lb scale weight.")
    print("  Residential delivery. Fuel included. Excludes additional handling, oversize,")
    print("  peak surcharges, address correction and Saturday delivery.")

    # ---- xlsx version ----
    wb = Workbook()
    ws = wb.active
    ws.title = "QUOTE"
    ws["A1"] = "SHIPO LLC — SHIPPING QUOTE"
    ws["A1"].font = Font(name=ARIAL, size=15, bold=True)
    ws["A2"] = "310 Cornell Dr, Wilmington, DE 19801  ·  302-442-2343  ·  shipousa.com"
    ws["A2"].font = Font(name=ARIAL, size=9)
    ws["A4"] = "Customer"
    ws["A4"].font = BOLD
    ws["B4"] = customer
    ws["B4"].font = BODY
    ws["A5"] = "Quote date"
    ws["A5"].font = BOLD
    ws["B5"] = str(data.get("pulledAt", ""))[:10]
    ws["B5"].font = BODY
    ws["C5"] = "Valid 14 days — UPS rates and fuel change."
    ws["C5"].font = Font(name=ARIAL, size=9, italic=True)

    heads = ["#", "To ZIP", "Service", "Zone", "Actual lb", "Billable lb", "PRICE"]
    for j, h in enumerate(heads):
        c = ws.cell(row=7, column=1 + j, value=h)
        c.font = BOLD
        c.fill = HDRFILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 8
    for i, l in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=i).font = BODY
        ws.cell(row=r, column=2, value=l["zip"]).font = BODY
        ws.cell(row=r, column=3, value=l["svc"]).font = BODY
        ws.cell(row=r, column=4, value=l["zone"]).font = BODY
        ws.cell(row=r, column=5, value=l["actual"]).font = BODY
        ws.cell(row=r, column=6, value=l["billable"]).font = BODY
        p = ws.cell(row=r, column=7, value=l["price"])
        p.font = BOLD
        p.number_format = MONEY
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = BOX
        r += 1
    t = ws.cell(row=r, column=7, value=f"=SUM(G8:G{r-1})")
    t.font = Font(name=ARIAL, size=12, bold=True)
    t.number_format = MONEY
    t.fill = HDRFILL
    t.border = BOX
    lbl = ws.cell(row=r, column=6, value="TOTAL")
    lbl.font = BOLD

    r += 2
    notes = [
        "Prices are per shipment, residential delivery, fuel included.",
        "Billable weight is the greater of actual weight and (L x W x H) / 166.",
        "Excludes: additional handling, oversize, peak-season surcharges,",
        "address corrections and Saturday delivery.",
        f"Rates quoted live from UPS via ShipStation on {str(data.get('pulledAt',''))[:10]}.",
    ]
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = Font(name=ARIAL, size=9)
        r += 1

    for col, wd in (("A", 5), ("B", 10), ("C", 26), ("D", 7), ("E", 11), ("F", 12), ("G", 12)):
        ws.column_dimensions[col].width = wd

    safe = "".join(ch if ch.isalnum() or ch in "-_ " else "" for ch in customer).strip().replace(" ", "-")
    path = os.path.join(OUT, f"Shipo-Quote-{safe or 'customer'}.xlsx")
    wb.save(path)
    print(f"\nwrote {path}")


if __name__ == "__main__":
    main()
